"""I: documents and presentations, against plugin/plugins/deliverables (spec checks, python-pptx deck build, HTML report
template, QA checks, output-path rules). Report PDFs: the laptop prints with Microsoft Edge; here the SAME HTML is printed
with headless Google Chrome (the same Chromium engine) as a stand-in, and the Edge print itself is a deferred case."""
from __future__ import annotations

import copy
import json
import os
import tempfile
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[5]
FIX = ROOT / "plugin/plugins/deliverables/tests/fixtures"
os.environ.setdefault("DELIVERABLES_OUT_DIR", tempfile.mkdtemp(prefix="neko-deliverables-"))
from plugin.plugins.deliverables import deck, qa, report, spec, tools  # noqa: E402


def _demo(name):
    return json.loads((FIX / f"demo_{name}.json").read_text(encoding="utf-8"))


def spec_ok(a):
    d, r = spec.deck(_demo("deck")), spec.report(_demo("report"))
    return len(d["slides"]) == 10 and len(r["sections"]) == 4, f"{len(d['slides'])} slides, {len(r['sections'])} sections"


def spec_refused(a):
    try:
        (spec.report if a.get("kind") == "report" else spec.deck)(a["spec"])
    except spec.SpecError as e:
        return a["msg"].lower() in str(e).lower(), f"refused: {e}"
    return False, "accepted"


def deck_builds(a):
    with tempfile.TemporaryDirectory() as t:
        out = Path(t) / "d.pptx"
        t0 = time.perf_counter()
        info = deck.build(spec.deck(_demo("deck")), str(out))
        secs = time.perf_counter() - t0
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            bad = z.testzip()
        charts = [n for n in names if n.startswith("ppt/charts/chart")]
        ok = info["slides"] == 11 and all(f["fits"] for f in info["fit"]) and bad is None and charts and secs < a.get("max_s", 10)
        return ok, f"{info['slides']} slides, all fit {all(f['fits'] for f in info['fit'])}, charts {len(charts)}, zip ok {bad is None}, {secs:.2f} s"


def deck_links_and_notes(a):
    with tempfile.TemporaryDirectory() as t:
        out = Path(t) / "d.pptx"
        deck.build(spec.deck(_demo("deck")), str(out))
        with zipfile.ZipFile(out) as z:
            rels = " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".rels"))
            notes = [n for n in z.namelist() if n.startswith("ppt/notesSlides/")]
        urls = [s["url"] for s in _demo("deck").get("sources", [])]
        linked = [u for u in urls if u.replace("&", "&amp;") in rels]
        return len(linked) == len(urls) and bool(notes), f"{len(linked)}/{len(urls)} source links clickable; {len(notes)} notes parts"


def overflow_reported(a):
    s = copy.deepcopy(_demo("deck"))
    s["slides"] = [{"layout": "bullets", "title": "Too much", "bullets": [("word " * 60).strip() for _ in range(a.get("n", 12))]}]
    try:
        sp = spec.deck(s)
    except spec.SpecError as e:
        return True, f"refused before building: {e}"
    with tempfile.TemporaryDirectory() as t:
        info = deck.build(sp, str(Path(t) / "o.pptx"))
    unfit = [f for f in info["fit"] if not f["fits"]]
    return bool(unfit), f"{len(unfit)} text boxes reported as not fitting (of {len(info['fit'])})"


def html_escapes(a):
    s = copy.deepcopy(_demo("report"))
    s["sections"][0]["paragraphs"] = ["<script>alert(1)</script><img src=x onerror=alert(2)> plain"]
    html, _t = report.render_html(spec.report(s))
    return "<script>alert(1)" not in html and "onerror=alert" not in html.replace("&lt;", "<").replace("onerror=alert(2)&gt;", "") \
        and "&lt;script&gt;" in html, "script and handler escaped"


def html_citations_link(a):
    html, _t = report.render_html(spec.report(_demo("report")))
    return 'href="#src-1"' in html or "#src" in html or "href=\"#S1" in html, "citation markers link to the source list"


def _print_pdf(html: str, path: Path, title: str = "Bench report"):
    """report.print_pdf exactly as the laptop runs it (JS off, network blocked, tagged PDF, bookmarks, footer), with only
    the browser channel swapped: msedge -> chrome."""
    import playwright.sync_api as PSA
    real = PSA.sync_playwright

    class _Swap:
        def __init__(self):
            self._cm = real()

        def __enter__(self):
            pw = self._cm.__enter__()
            launch = pw.chromium.launch

            def chrome_launch(**kw):
                if kw.get("channel") == "msedge":
                    kw["channel"] = "chrome"
                return launch(**kw)
            pw.chromium.launch = chrome_launch
            return pw

        def __exit__(self, *e):
            return self._cm.__exit__(*e)
    PSA.sync_playwright = _Swap
    try:
        report.print_pdf(html, str(path), title)
    finally:
        PSA.sync_playwright = real


def report_pdf_qa(a):
    sp = spec.report(_demo("report"))
    html, _t = report.render_html(sp)
    with tempfile.TemporaryDirectory() as t:
        pdf = Path(t) / "r.pdf"
        _print_pdf(html, pdf)
        off = qa.words_off_page(pdf)
        res = qa.check_report(pdf, sp, Path(t))
        failed = [k for k, v in res["checks"].items() if not v]
        return not off and res.get("passed"), f"words off page {len(off)}; QA passed {res.get('passed')}; failed checks {failed}; warnings {res['warnings'][:2]}"


def outputs_only(a):
    root = Path(os.environ["DELIVERABLES_OUT_DIR"])
    tools.OUT_ROOT = root
    good = tools.new_job_dir("Quarterly Report: Q3/2026 <draft>")
    f = good / "x.pdf"
    f.write_bytes(b"%PDF-1.4\n")
    bad = []
    for path in (str(root / ".." / "etc" / "passwd"), "/etc/hostname", str(good / "x.exe"), str(good / ".." / ".." / "x.pdf")):
        try:
            tools.inside_outputs(path)
            bad.append(path)
        except tools.ToolError:
            pass
    ok_open = tools.inside_outputs(str(f)) == f.resolve()
    return not bad and ok_open and good.parent == root, f"accepted outside paths {bad}; own file ok {ok_open}; folder {good.name}"


def unique_job_dirs(a):
    tools.OUT_ROOT = Path(os.environ["DELIVERABLES_OUT_DIR"])
    ds = [tools.new_job_dir("Same title") for _ in range(3)]
    return len({d for d in ds}) == 3, f"{[d.name for d in ds]}"


def missing_edge_honest(a):
    import plugin.plugins.deliverables.report as R
    try:
        R.print_pdf("<p>x</p>", os.path.join(tempfile.mkdtemp(), "x.pdf"), "x")
    except tools.ToolError as e:
        return "Edge" in str(e) and "install" not in str(e).lower(), f"ToolError: {e}"
    except Exception as e:  # noqa: BLE001
        return False, f"raw {type(e).__name__}: {str(e)[:120]}"
    return False, "printed (Edge present?)"


def missing_officecli_honest(a):
    with tempfile.TemporaryDirectory() as t:
        out = Path(t) / "d.pptx"
        deck.build(spec.deck(_demo("deck")), str(out))
        old = tools.OFFICECLI
        tools.OFFICECLI = "/nonexistent/officecli.exe"
        try:
            from plugin.plugins.deliverables import animate
            animate.apply(out, 11, "fade", [], {1})
            return False, "animation claimed success without OfficeCLI"
        except tools.ToolError as e:
            return True, f"ToolError: {str(e)[:120]}"
        except Exception as e:  # noqa: BLE001
            return False, f"raw {type(e).__name__}: {str(e)[:120]}"
        finally:
            tools.OFFICECLI = old


# ================================================================ added 2026-10-17: untrusted specs, malformed input and output
import re  # noqa: E402
import threading  # noqa: E402


def _refused(fn, s):
    """'SpecError: ...', 'ACCEPTED', or 'RAW <type>' for anything else (a crash the plugin entry does not catch)."""
    try:
        fn(s)
        return "ACCEPTED"
    except spec.SpecError as e:
        return "SpecError: " + str(e)[:60]
    except Exception as e:  # noqa: BLE001
        return f"RAW {type(e).__name__}"


def spec_wrong_types(a):
    """A model's spec with the right keys but wrong types (a number as layout, chart type or section kind; a word as a
    bullet level) is refused with a clear SpecError, never a raw crash."""
    deck_bad = {"layout as number": {"title": "T", "slides": [{"layout": 5, "title": "x"}]},
                "level as word": {"title": "T", "slides": [{"layout": "bullets", "title": "x", "bullets": [{"text": "a", "level": "high"}]}]},
                "chart type as number": {"title": "T", "slides": [{"layout": "chart", "title": "x", "chart": {"type": 3, "labels": ["a"], "series": [{"values": [1]}]}}]}}
    rep_bad = {"kind as number": {"title": "T", "sections": [{"heading": "h", "kind": 3}]},
               "level as word": {"title": "T", "sections": [{"heading": "h", "bullets": [{"text": "a", "level": "x"}]}]}}
    got = {n: _refused(spec.deck, s) for n, s in deck_bad.items()}
    got.update({"report " + n: _refused(spec.report, s) for n, s in rep_bad.items()})
    bad = {n: g for n, g in got.items() if not g.startswith("SpecError")}
    return not bad, f"not a clear refusal: {bad}" if bad else f"all {len(got)} refused with SpecError"


def chart_nonfinite_refused(a):
    """'nan', 'inf' and 1e309 as chart values are numbers to float() but cannot be drawn or saved: refused before any
    file is made (they used to pass the spec and fail inside the chart writer)."""
    got = {}
    for vals in (["nan", 1], ["inf", 1], [1e309, 1], ["-inf", 1]):
        s = {"title": "T", "slides": [{"layout": "chart", "title": "c", "chart": {"type": "column", "labels": ["a", "b"], "series": [{"name": "s", "values": vals}]}}]}
        got[repr(vals[0])] = _refused(spec.deck, s)
    bad = {k: v for k, v in got.items() if not v.startswith("SpecError")}
    return not bad, f"accepted or crashed: {bad}" if bad else "all refused"


def control_chars_in_deck(a):
    """Text extracted from documents often carries control characters (vertical tab, form feed, NUL, bell). The deck
    builds and reopens, and the audience never sees codes like _x000B_ on a slide or in the notes."""
    s = copy.deepcopy(_demo("deck"))
    s["slides"] = [{"layout": "bullets", "title": "Extracted\x0btext", "bullets": ["line one\x0cline two", "nul\x00here", "bell\x07sound"], "notes": "speaker\x01notes"}]
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as t:
        out = Path(t) / "c.pptx"
        try:
            deck.build(spec.deck(s), str(out))
            prs = Presentation(str(out))
        except Exception as e:  # noqa: BLE001
            return False, f"{type(e).__name__}: {str(e)[:120]}"
        texts = [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]
        texts += [sl.notes_slide.notes_text_frame.text for sl in prs.slides if sl.has_notes_slide]
    codes = sorted({m for x in texts for m in re.findall(r"_x00[0-9A-F]{2}_", x)})
    words = all(any(w in x for x in texts) for w in ("Extracted", "line two", "here", "sound", "notes"))
    return not codes and words, f"codes shown to the audience: {codes or 'none'}; words kept={words}"


def report_no_remote_images(a):
    """Untrusted report text with markdown links to javascript:, file:, data: and vbscript: and an image pointing at an
    internal address: the report HTML (also saved next to the PDF and openable) keeps only web and in-page links and
    loads nothing from elsewhere."""
    r = copy.deepcopy(_demo("report"))
    r["sections"][0]["paragraphs"] = ["See [a](javascript:alert(1)), [b](file:///etc/passwd), [c](data:text/html,x), [d](vbscript:x) and "
                                      "![pixel](http://127.0.0.1:9/track.png?u=owner) and [ok](https://ok.example/x)."]
    r["sections"][0]["callout"] = "![beacon](https://tracker.example/p.gif)"
    html, _t = report.render_html(spec.report(r))
    hrefs = re.findall(r'href="([^"]*)"', html)
    bad_links = [h for h in hrefs if not h.startswith(("https://", "http://", "#"))]
    remote = re.findall(r'<(?:img|iframe|script|link|object|embed)\b[^>]*(?:src|href|data)="(?:https?:)?//[^"]*"', html)
    return not bad_links and not remote, f"non-web links: {bad_links or 'none'}; remote loads: {remote or 'none'}"


class _Hits:
    def __init__(self):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        hits = self.hits = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *x):
                pass

            def do_GET(self):
                hits.append(self.path)
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Content-Length", "0")
                self.end_headers()
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"

    def close(self):
        self.srv.shutdown()
        self.srv.server_close()


def report_print_offline(a):
    """Printing a report whose text embeds an image from a local server: the print (JS off, every request blocked; Chrome
    stand-in for Edge) makes no request at all, and the PDF is made."""
    srv = _Hits()
    try:
        r = copy.deepcopy(_demo("report"))
        r["sections"][0]["paragraphs"] = [f"Status ![pixel]({srv.base}/track.png?u=owner) here.", f"[link]({srv.base}/page)"]
        html, _t = report.render_html(spec.report(r))
        with tempfile.TemporaryDirectory() as t:
            pdf = Path(t) / "r.pdf"
            _print_pdf(html, pdf)
            made = pdf.is_file() and pdf.stat().st_size > 1000
    finally:
        srv.close()
    return made and not srv.hits, f"PDF made={made}; requests to the local server={srv.hits}"


def report_attribute_breakout(a):
    """A source URL with '\" onmouseover=\"...' and a title, heading and table cell with HTML: nothing becomes an element
    or an attribute in the report HTML."""
    from lxml import html as lhtml
    r = copy.deepcopy(_demo("report"))
    r["sources"][0]["url"] = 'https://x.example/a" onmouseover="alert(1)'
    r["title"] = 'Q3 </title><script>alert(1)</script>'
    r["sections"][0]["heading"] = '<img src=x onerror=alert(2)>Heading'
    r["sections"][0]["table"] = {"columns": ["a"], "rows": [['<iframe src="https://evil.example/"></iframe>']]}
    html, _t = report.render_html(spec.report(r))
    doc = lhtml.fromstring(html)
    handlers = [k for el in doc.iter() if isinstance(el.tag, str) for k in el.attrib if k.lower().startswith("on")]
    bad = [el.tag for el in doc.iter() if el.tag in ("script", "iframe", "img")]
    return not handlers and not bad, f"event attributes: {handlers or 'none'}; injected elements: {bad or 'none'}"


def footer_title_escaped(a):
    """The report title goes into the printed page footer (Chrome stand-in): markup in it is printed as text, not run."""
    r = copy.deepcopy(_demo("report"))
    r["title"] = "Plan <b>bold</b> <img src=x>"
    sp = spec.report(r)
    html, _t = report.render_html(sp)
    import pdfplumber
    with tempfile.TemporaryDirectory() as t:
        pdf = Path(t) / "f.pdf"
        _print_pdf(html, pdf, sp["title"])
        with pdfplumber.open(str(pdf)) as doc:
            last = doc.pages[-1].extract_text() or ""
    return "<b>bold</b>" in last and "<img src=x>" in last, f"footer text on the last page: {last.splitlines()[-1][:120] if last else ''!r}"


def outputs_symlink_refused(a):
    """A link inside the outputs folder that points outside it (as a file named .pdf, or as a folder) cannot be opened."""
    root = Path(tempfile.mkdtemp(prefix="neko-deliv-link-"))
    outside = Path(tempfile.mkdtemp(prefix="neko-outside-"))
    tools.OUT_ROOT = root
    try:
        (outside / "secret.pdf").write_bytes(b"%PDF-1.4 owner data\n")
        job = tools.new_job_dir("Linked")
        os.symlink(outside / "secret.pdf", job / "report.pdf")
        os.symlink(outside, job / "more")
        bad = []
        for p in (job / "report.pdf", job / "more" / "secret.pdf"):
            try:
                tools.inside_outputs(str(p))
                bad.append(str(p.relative_to(root)))
            except tools.ToolError:
                pass
    finally:
        import shutil
        shutil.rmtree(root, ignore_errors=True)
        shutil.rmtree(outside, ignore_errors=True)
        tools.OUT_ROOT = Path(os.environ["DELIVERABLES_OUT_DIR"])
    return not bad, f"opened through a link: {bad or 'none'}"


def regenerate_outside_refused(a):
    """regenerate() reads a spec.json only from a deliverable folder: a folder outside the outputs root, or a link inside
    it pointing outside, is refused before anything is read or built."""
    from plugin.plugins.deliverables import core as dcore
    root = Path(tempfile.mkdtemp(prefix="neko-deliv-regen-"))
    outside = Path(tempfile.mkdtemp(prefix="neko-outside-"))
    tools.OUT_ROOT = root
    built = []
    real = dcore.make_report
    dcore.make_report = lambda *x, **k: built.append(1) or {}
    try:
        (outside / "spec.json").write_text(json.dumps(_demo("report")), encoding="utf-8")
        os.symlink(outside, root / "20260101-000000-linked")
        errs = []
        for f in (outside, root / "20260101-000000-linked"):
            try:
                dcore.regenerate(str(f))
                errs.append("accepted")
            except tools.ToolError:
                errs.append("refused")
    finally:
        dcore.make_report = real
        import shutil
        shutil.rmtree(root, ignore_errors=True)
        shutil.rmtree(outside, ignore_errors=True)
        tools.OUT_ROOT = Path(os.environ["DELIVERABLES_OUT_DIR"])
    return errs == ["refused", "refused"] and not built, f"outside folder / linked folder -> {errs}; builds started={len(built)}"


def _pdf(path, words, outside=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(4, 3))
    fig.text(0.1, 0.5, words)
    if outside:
        fig.text(1.3, 0.5, outside, clip_on=False)
    fig.savefig(path)
    plt.close(fig)


def qa_off_page_text(a):
    """QA reads a PDF the way a careful person would: words drawn beyond the page edge are reported; a clean page has none."""
    with tempfile.TemporaryDirectory() as t:
        bad, good = Path(t) / "bad.pdf", Path(t) / "good.pdf"
        _pdf(bad, "Inside words here", outside="OVERFLOW words")
        _pdf(good, "Inside words here")
        off, clean = qa.words_off_page(bad), qa.words_off_page(good)
    return any(w["text"] == "OVERFLOW" for w in off) and not clean, f"off-page words found={off}; on the clean page={clean}"


def qa_damaged_pdf(a):
    """A PDF cut in half, and random bytes behind a %PDF header: QA reports 'not passed' with a reason; it never says
    passed and never fails with a raw parser error."""
    sp = spec.report(_demo("report"))
    out = {}
    with tempfile.TemporaryDirectory() as t:
        good = Path(t) / "g.pdf"
        _pdf(good, "Some report text")
        data = good.read_bytes()
        for name, body in (("cut", data[: len(data) // 2]), ("junk", b"%PDF-1.7\n" + bytes((i * 91 + 7) % 256 for i in range(4000)))):
            p = Path(t) / f"{name}.pdf"
            p.write_bytes(body)
            try:
                res = qa.check_report(p, sp, Path(t))
                out[name] = "passed" if res.get("passed") else "not passed"
            except Exception as e:  # noqa: BLE001
                out[name] = f"RAW {type(e).__name__}"
    return all(v == "not passed" for v in out.values()), f"{out}"


def deck_only_source_links(a):
    """Bullets and notes carry a markdown javascript: link, an HTML anchor to file:, and a bare phishing URL: in the
    built deck the only hyperlinks are the listed sources' web addresses."""
    s = copy.deepcopy(_demo("deck"))
    for sl in s["slides"]:
        if sl.get("layout") == "bullets":
            sl["bullets"] = ["[click me](javascript:alert(1)) and <a href='file:///etc/passwd'>files</a>", "visit https://evil.example/phish now"]
            sl["notes"] = "[notes link](javascript:steal())"
    allowed = {x["url"] for x in s["sources"]}
    with tempfile.TemporaryDirectory() as t:
        out = Path(t) / "l.pptx"
        deck.build(spec.deck(s), str(out))
        with zipfile.ZipFile(out) as z:
            rels = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".rels"))
    targets = set(re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', rels)) | set(re.findall(r'TargetMode="External"[^>]*Target="([^"]+)"', rels))
    targets = {x.replace("&amp;", "&") for x in targets}
    extra = sorted(targets - allowed)
    return not extra and targets, f"external links not from the source list: {extra or 'none'}; links={len(targets)}"


def research_deck_keeps_citations(a):
    """A checked research report turned into a deck (research core.to_deck): long statements are cut to fit a slide but
    keep every citation marker, and the deck spec is valid."""
    from plugin.plugins.research import core as rcore
    long_ = ("Residential solar panels convert sunlight into electricity with silicon cells, and their output depends on "
             "panel size, the cell type, the roof angle, local weather, shading from trees and the temperature of the cells. [S1][S2]")
    summary = ("Across the sources, a typical residential panel produces between 250 and 400 watts in full sun, most panels keep "
               "about 80 percent of that output after 25 years of use, and installers size a home system from the yearly electricity "
               "bill, the usable roof area facing the sun, local net-metering rules, the cost of batteries and the expected rise in "
               "electricity prices over the next decade, which together decide how many panels pay for themselves. [S1][S2]")
    sp = {"title": "Solar", "subtitle": "", "author": "Neko research", "topic": "solar", "summary": summary,
          "sources": [{"id": "S1", "title": "A", "url": "https://a.example/", "date": "", "publisher": ""},
                      {"id": "S2", "title": "B", "url": "https://b.example/", "date": "", "publisher": ""}],
          "sections": [{"heading": "Output", "kind": "fact", "paragraphs": [long_], "bullets": [long_ + " Extra words here. [S2]"]}]}
    try:
        d = spec.deck(rcore.to_deck(sp))
    except spec.SpecError as e:
        return False, f"the deck spec made from the report was refused: {e}"
    items = [b["text"] for sl in d["slides"] for b in sl["bullets"]]
    cut = [x for x in items if "…" in x]
    lost = [x[:50] for x in items if ("silicon cells" in x or "net-metering" in x or "250 and 400" in x) and not ("[S1]" in x and "[S2]" in x)]
    return len(cut) >= 2 and not lost, f"cut items={len(cut)}; items that lost a citation={lost or 'none'}"


def job_dirs_concurrent(a):
    """Eight deliverables with the same title started at the same moment each get their own new folder (no error, no
    shared folder)."""
    tools.OUT_ROOT = Path(tempfile.mkdtemp(prefix="neko-deliv-race-"))
    errs, dirs = [], []
    try:
        for _ in range(5):
            go = threading.Barrier(8)

            def mk():
                go.wait()
                try:
                    dirs.append(tools.new_job_dir("Same title"))
                except Exception as e:  # noqa: BLE001
                    errs.append(type(e).__name__)
            ts = [threading.Thread(target=mk) for _ in range(8)]
            [x.start() for x in ts]
            [x.join() for x in ts]
    finally:
        tools.OUT_ROOT = Path(os.environ["DELIVERABLES_OUT_DIR"])
    return not errs and len(set(dirs)) == 40, f"errors={sorted(set(errs)) or 'none'} ({len(errs)}); distinct folders={len(set(dirs))}/40"


def deck_max_size(a):
    """The largest deck a spec allows (40 slides of 12 bullets) builds within the time limit and reports the fit of
    every text box."""
    s = copy.deepcopy(_demo("deck"))
    s["slides"] = [{"layout": "bullets", "title": f"Slide {i + 1}", "bullets": [f"Point {j + 1} about panel output, roof angle and shading" for j in range(12)],
                    "notes": "Speaker notes " * 20} for i in range(40)]
    with tempfile.TemporaryDirectory() as t:
        t0 = time.perf_counter()
        info = deck.build(spec.deck(s), str(Path(t) / "max.pptx"))
        secs = time.perf_counter() - t0
        with zipfile.ZipFile(Path(t) / "max.pptx") as z:
            intact = z.testzip() is None
    slides_fit = {f["slide"] for f in info["fit"]}
    ok = info["slides"] >= 40 and intact and len(slides_fit) == info["slides"] and secs < a["max_s"]
    return ok, f"{info['slides']} slides in {secs:.1f}s (limit {a['max_s']}); fit reported for {len(slides_fit)} slides; " \
               f"not fitting={sum(not f['fits'] for f in info['fit'])}; zip ok={intact}"
