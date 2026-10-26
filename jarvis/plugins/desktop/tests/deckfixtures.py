"""Shared pieces of the D64 playbook and deck-flow tests: a real .pptx and .pdf made on the fly, and the playbooks of the two test sites."""
from __future__ import annotations

import io

from pptx import Presentation


def make_pptx(slides: int = 3) -> bytes:
    prs = Presentation()
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide {i + 1} title"
        s.placeholders[1].text = "A line of body text for the test deck."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf(pages: int = 3) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=300, height=200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def decks_playbook(host: str = "decks.test") -> dict:
    return {
        "name": "make_deck", "host": host, "hosts": [host], "version": 1, "source": "builtin", "goal": "Make a deck",
        "start": f"https://{host}/create", "verified_live": False,
        "needs": ["signin_existing_session", "accept_terms", "download_to_quarantine"],
        "terms_note": "Test site: no automation clause.", "terms_read": "2026-10-26", "terms_source": "decks.test/terms",
        "formats": ["pptx", "pdf"],
        "params": [{"name": "topic", "required": True}, {"name": "slides", "default": "3"}, {"name": "theme", "default": ""}, {"name": "formats", "default": "pptx"}],
        "phases": [
            {"name": "open", "steps": [{"id": "open", "do": "goto", "to": f"https://{host}/create", "new_tab": True, "expect": {"host_is": host}}]},
            {"name": "signin", "skip_if": {"element": {"role": "textbox", "name": "What is your presentation about?"}},
             "steps": [
                 {"id": "google", "do": "click", "loc": {"role": "button", "name": "Continue with Google"}, "expect": {"any": [{"host_is": "accounts.google.com"}, {"host_is": host}]}},
                 {"id": "account", "do": "click", "loc": {"role": "button", "text": "@"}, "optional": True},
                 {"id": "consent", "do": "click", "loc": {"role": "button", "name": "Continue"}, "optional": True, "expect": {"host_is": host}}]},
            {"name": "create", "steps": [
                {"id": "prompt", "do": "fill", "loc": {"role": "textbox", "name": "What is your presentation about?"}, "value": "{topic}"},
                {"id": "theme", "do": "select", "loc": {"role": "combobox", "name": "Theme"}, "value": "{theme}", "optional": True},
                {"id": "generate", "do": "click", "loc": {"role": "button", "name": "Generate"}, "expect": {"text": "Generating slide"}}]},
            {"name": "wait", "steps": [
                {"id": "ready", "do": "wait", "timeout_s": 60, "until": {"text": "Your deck is ready", "fail_if": ["Something went wrong"]}}]},
            {"name": "export", "steps": [
                {"id": "export", "do": "click", "loc": {"role": "button", "name": "Export"}, "expect": {"text": "Export your deck"}},
                {"id": "pptx_go", "do": "click", "loc": {"role": "button", "name": "Download PowerPoint"}, "expect": {"noop_ok": True}},
                {"id": "pptx_file", "do": "download", "value": "pptx", "expect": {"download": {"ext": "pptx"}}}]},
        ]}


def second_playbook() -> dict:
    return {
        "name": "make_deck", "host": "second.test", "hosts": ["second.test"], "version": 1, "source": "builtin", "goal": "Make a deck",
        "start": "https://second.test/", "verified_live": False, "needs": ["download_to_quarantine"], "terms_note": "Test site.", "formats": ["pptx"],
        "params": [{"name": "topic", "required": True}, {"name": "slides", "default": "3"}, {"name": "theme", "default": ""}, {"name": "formats", "default": "pptx"}],
        "phases": [
            {"name": "open", "steps": [{"id": "open", "do": "goto", "to": "https://second.test/", "new_tab": True, "expect": {"host_is": "second.test"}}]},
            {"name": "create", "steps": [
                {"id": "prompt", "do": "fill", "loc": {"role": "textbox", "name": "Topic"}, "value": "{topic}"},
                {"id": "make", "do": "click", "loc": {"role": "button", "name": "Make slides"}, "expect": {"text": "Slides ready"}}]},
            {"name": "export", "steps": [
                {"id": "get", "do": "click", "loc": {"role": "button", "name": "Get PowerPoint"}, "expect": {"noop_ok": True}},
                {"id": "file", "do": "download", "value": "pptx", "expect": {"download": {"ext": "pptx"}}}]},
        ]}


def shortlist(*rows) -> dict:
    return {"checked": "2026-10-26", "sites": [dict(r) for r in rows]}


def row(host: str, rank: int, status: str = "shortlist") -> dict:
    return {"id": host.split(".")[0], "name": host, "host": host, "hosts": [host], "status": status, "rank": rank}
