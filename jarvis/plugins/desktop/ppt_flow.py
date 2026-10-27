"""The free presentation generator flow (D64, 2026-10-26, owner: "it can go to a FREE PPT generator website, log in by itself, and generate everything").

JARVIS does not build the deck. It opens a free AI presentation generator in the owner's own Chrome, uses his EXISTING signed-in session (or follows the sign-in
policy below), gives the site the topic (or an outline the online model wrote first, text only, through the privacy gateway), picks a theme, generates, waits for
completion BY CONDITION, edits slide text if asked, exports to the QUARANTINE folder, verifies the file (it exists, its type, its size, its slide count read with the
document readers), shows it in the Preview panel and reports only what was verified. When the first site fails or hits its free limit, the next one on the shortlist is
tried automatically; only when all have failed is the owner asked, with a form that shows what blocked each one.

The shortlist is data (config/ppt_sites.json: one row per site with the vendor's own words, their source and the date read) and each site has a playbook
(config/site_playbooks/<host>.make_deck.json), which the laptop must still confirm against the real site (`verified_live`).

Sign-in policy (kernel: sitegrants.py, policy.py). The kernel treats sign-in, accepting terms, signing up and downloading as class C. A ONE-TIME PER-SITE GRANT, which
the owner confirms with the PIN in the console, lets JARVIS do exactly those things on exactly that host until it expires. This module only ASKS for the card
(`GrantsClient.request`) and then waits for the PIN; it never grants anything, and a site without a grant is never driven past the point that needs one.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import time
import urllib.parse
from dataclasses import dataclass, field
from pathlib import Path

from . import browser_agent as BA
from . import browser_playbooks as PB
from . import browser_snapshot as S

JARVIS = Path(__file__).resolve().parents[2]
SHORTLIST_FILE = JARVIS / "config" / "ppt_sites.json"
BANNERS_FILE = JARVIS / "config" / "cookie_banners.json"
OUTLINE_MAX = 470
OUTLINE_SYSTEM = ("You write the outline of a short slide deck for a presentation generator website. Reply with plain text only: a title line, then one line per slide "
                  "('Slide N: heading - one sentence'), 450 characters in all. No markdown, no commentary, no links.")


class FlowError(Exception):
    pass


# ---------------------------------------------------------------------------------------------------------------------- data files
def load_shortlist(path: Path = None) -> dict:
    try:
        data = json.loads(Path(path or SHORTLIST_FILE).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return {"checked": "", "sites": []}
    sites = [s for s in data.get("sites", []) if isinstance(s, dict) and s.get("host")]
    return {**data, "sites": sites}


def ranked_sites(shortlist: dict = None, only=None) -> list:
    """The sites to try, best first: the ranked shortlist, then the fallbacks. Excluded sites (terms that forbid automation, a block, no free plan) are never tried."""
    sl = shortlist or load_shortlist()
    rows = [s for s in sl["sites"] if s.get("status") in ("shortlist", "fallback")]
    if only:
        rows = [s for s in rows if s["host"] in only]
    return sorted(rows, key=lambda s: (0 if s.get("status") == "shortlist" else 1, int(s.get("rank", 99))))


def load_banners(path: Path = None) -> dict:
    try:
        d = json.loads(Path(path or BANNERS_FILE).read_text(encoding="utf-8-sig"))
        return {"reject": [str(x) for x in d.get("reject", [])], "accept": [str(x) for x in d.get("accept", [])]}
    except (OSError, ValueError):
        return {"reject": ["Reject all", "Reject", "Necessary only", "Only necessary"], "accept": []}


# ---------------------------------------------------------------------------------------------------------------------- the file the site gave us
def verify_file(path: str) -> dict:
    """What is TRUE about the downloaded file, read from the file itself (python-pptx for a deck, pypdf for a PDF): its type, size, sha256 and its slide or page
    count. Raises FlowError when it cannot be read as what its name says."""
    p = Path(path)
    if not p.is_file():
        raise FlowError("the file is not on disk")
    ext = p.suffix.lower().lstrip(".")
    info = {"name": p.name, "ext": ext, "size": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()[:16]}
    if info["size"] < 200:
        raise FlowError("the file is empty or too small to be a document")
    try:
        if ext == "pptx":
            from pptx import Presentation
            prs = Presentation(str(p))
            titles = []
            for sl in prs.slides:
                t = ""
                for sh in sl.shapes:
                    if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                        t = sh.text_frame.text.strip().splitlines()[0][:80]
                        break
                titles.append(t)
            info.update(slides=len(prs.slides), titles=titles)
        elif ext == "pdf":
            from pypdf import PdfReader
            r = PdfReader(str(p))
            info.update(pages=len(r.pages))
        elif ext in ("png", "jpg", "jpeg"):
            info.update(image=True)
    except Exception as e:  # noqa: BLE001
        raise FlowError(f"the {ext} file could not be read ({type(e).__name__})") from None
    if info.get("slides") == 0 or info.get("pages") == 0:
        raise FlowError("the file has no slides")
    return info


def promote(path: str, topic: str, deliverables_dir: str = None, now=time.time) -> str:
    """A verified file is COPIED (the quarantine copy stays) into the deliverables folder under a plain name, so the Preview panel can show it. Returns the new path."""
    root = Path(deliverables_dir or os.environ.get("DELIVERABLES_OUT_DIR", r"D:\jarvis\outputs\deliverables"))
    root.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"[^a-z0-9]+", "-", topic.lower()).strip("-")[:40] or "deck"
    dest = root / (time.strftime("%Y%m%d-%H%M%S-", time.localtime(now())) + "web-" + slug + Path(path).suffix.lower())
    n = 1
    while dest.exists():
        n += 1
        dest = dest.with_name(dest.stem + f"-{n}" + dest.suffix)
    shutil.copy2(path, dest)
    return str(dest)


# ---------------------------------------------------------------------------------------------------------------------- account creation (only with the owner's 'signup' grant)
GENERIC_SIGNUP = "signup_generic"


def signup_phase() -> PB.Phase:
    """The generic steps of creating an account on a site the owner granted 'signup' for: the alias e-mail from config/identity.json and a generated password, both
    typed BY THE NODE from references (@identity:email, @cred:password); the terms box ticked (accept_terms is part of the same grant); then the submit. Found by
    meaning (role + name + synonyms), like every other step. E-mail verification is the owner's: the engine meets the 'check your e-mail' wall and asks him."""
    def L(role, name, *more):
        return S.Locator(role=role, name=name, any_of=tuple(more))
    steps = [
        BA.Step("open_signup", "click", L("link", "Sign up", "Create account", "Create an account", "Get started", "Register", "Sign up for free", "Join"), optional=True,
                intent="open the account creation form", password_ok=True),
        BA.Step("pw_made", "cred_new", password_ok=True, intent="make the password for this site (kept by the node, never shown)"),
        BA.Step("su_name", "fill", L("textbox", "Full name", "Name", "Your name", "First name"), value="@identity:name", optional=True, password_ok=True, intent="the display name"),
        BA.Step("su_email", "fill", L("textbox", "Email", "Email address", "E-mail", "Work email"), value="@identity:email", password_ok=True, intent="the alias e-mail"),
        BA.Step("su_pw", "fill", L("textbox", "Password", "Create password", "Choose a password"), value="@cred:password", password_ok=True, intent="the generated password"),
        BA.Step("su_pw2", "fill", L("textbox", "Confirm password", "Repeat password", "Re-enter password"), value="@cred:password", optional=True, password_ok=True,
                intent="the generated password again"),
        BA.Step("su_terms", "check", L("checkbox", "I agree to the Terms of Service", "I agree", "Terms of Service", "Terms", "Privacy Policy"), optional=True, password_ok=True,
                intent="the terms box"),
        BA.Step("su_go", "click", L("button", "Create account", "Sign up", "Register", "Get started", "Continue", "Create"), password_ok=True, intent="create the account",
                expect={"changed": True})]
    return PB.Phase(GENERIC_SIGNUP, steps, {}, "account creation with the alias e-mail and a generated password", fallback=True)


def with_signup(pb: PB.Playbook) -> PB.Playbook:
    """A copy of the playbook that tries account creation when the sign-in phase could not be done. Not saved with the playbook."""
    c = PB.Playbook.from_dict(pb.to_dict(), pb.source)
    if any(p.fallback for p in c.phases):
        return c
    at = next((i for i, p in enumerate(c.phases) if p.name == "signin"), None)
    if at is None:
        return c
    c.phases.insert(at + 1, signup_phase())
    return c


# ---------------------------------------------------------------------------------------------------------------------- the outline
def default_outline(topic: str, slides: int, ask_online=None) -> "tuple[str, str] | None":
    """An outline written by the online model through the privacy gateway: TEXT ONLY, from the owner's own topic, with no page text in it. (text, provider) or None."""
    try:
        if ask_online is None:
            from .goal import ask_online as ask_online
        text, provider = ask_online([{"role": "system", "content": OUTLINE_SYSTEM},
                                     {"role": "user", "content": f"Topic: {topic[:300]}\nSlides: {slides}"}], max_tokens=500)
    except Exception:  # noqa: BLE001 - no outline is not a failure: the topic alone is given to the site
        return None
    text = re.sub(r"[\r\x00-\x08\x0b-\x1f]", "", str(text)).strip()
    return (text[:1500], provider) if text else None


# ---------------------------------------------------------------------------------------------------------------------- the result
@dataclass
class SiteAttempt:
    host: str
    outcome: str                  # done | failed | limit | needs_owner | no_playbook | no_grant | stopped
    why: str = ""
    rungs: list = field(default_factory=list)
    repaired: int = 0


@dataclass
class DeckResult:
    ok: bool
    site: str = ""
    files: list = field(default_factory=list)          # verified facts per file (no paths)
    facts: list = field(default_factory=list)          # sentences that were checked
    tried: list = field(default_factory=list)
    needs_owner: "dict | None" = None
    shown: str = ""
    local_paths: list = field(default_factory=list)    # on this laptop only; never in a model-facing answer
    summary: str = ""


# ---------------------------------------------------------------------------------------------------------------------- the flow
class DeckFlow:
    """`agent_factory(hosts) -> BA.Agent` makes a fresh engine (and backend) for one site; `grants` is browser_backend.GrantsClient; `ask` is the owner form."""

    def __init__(self, agent_factory, store: PB.Store, grants, shortlist: dict = None, outline_fn=None, show=None, take_local=None, stop=None,
                 grant_wait_s: float = 300.0, deliverables_dir: str = None, journal=None, ask=None, want_signup: bool = False):
        self.agent_factory, self.store, self.grants = agent_factory, store, grants
        self.ask, self.want_signup = ask, want_signup
        self.shortlist = shortlist or load_shortlist()
        self.outline_fn = outline_fn if outline_fn is not None else default_outline
        self.show, self.take_local, self.stop = show, take_local, stop or (lambda: False)
        self.grant_wait_s, self.deliverables_dir, self.journal = grant_wait_s, deliverables_dir, journal

    # ---- grants -------------------------------------------------------------------------------------------------------
    def ensure_grants(self, pb: PB.Playbook, signup: bool = False) -> None:
        """The grant card for this site, if it lacks one. The owner confirms it with the PIN in the console; this waits for him (bounded). No grant, no drive."""
        need = list(dict.fromkeys(pb.needs + (["signup"] if signup else [])))
        host = pb.hosts[0]
        lacking = [h for h in pb.hosts[:BA.MAX_GRANT_HOSTS] if self.grants.missing(h, need)]
        if not lacking:
            return
        req = {"hosts": lacking, "abilities": need, "days": 30, "note": pb.terms_note}
        got = self.grants.request(req["hosts"], req["abilities"], req["days"], req["note"])
        if "approval_id" not in got:
            raise BA.NeedOwner("grant", got.get("error", "the grant card could not be made"), {"site": host})
        done = self.grants.wait_and_apply(got["request"], got["approval_id"], self.grant_wait_s, self.stop)
        if "granted" not in done:
            raise BA.NeedOwner("grant", f"waiting for your PIN in the JARVIS console to let JARVIS act on {host} "
                                        f"({', '.join(need)}); approve the card, then ask me to continue", {"site": host, "approval_id": got["approval_id"]})

    # ---- one site -----------------------------------------------------------------------------------------------------
    def run_site(self, row: dict, topic: str, slides: int, theme: str, formats, outline: str, edits: list, result: DeckResult):
        host = row["host"]
        pb = self.store.get(host, "make_deck")
        if pb is None:
            result.tried.append(SiteAttempt(host, "no_playbook", "there is no playbook for this site yet"))
            return None
        try:
            self.ensure_grants(pb, signup=self.want_signup)
        except BA.NeedOwner as e:
            result.tried.append(SiteAttempt(host, "no_grant", e.detail))
            result.needs_owner = {"kind": "grant", "site": host, "detail": e.detail}
            raise
        if self.grants.has(host, "signup"):
            pb = with_signup(pb)
        agent = self.agent_factory(pb.hosts)
        agent.signin_site = host
        agent.goal = f"make a {slides}-slide deck about the owner's topic"
        runner = PB.Runner(agent, self.store)
        slides = min(slides, int(row.get("max_slides") or slides))           # the site's own free cap (Gamma: 10 slides in a single prompt); the report says what was asked
        params = {"topic": outline or topic, "slides": str(slides), "theme": theme or pb.defaults().get("theme", ""), "formats": ",".join(formats), "outline": outline or topic}
        files: list = []
        try:
            r = runner.run(pb, params, phases=None)
            for st_res in r.results:
                ref = st_res.data.get("local_result") if st_res.data else None
                if ref and self.take_local:
                    doc = self.take_local(ref)
                    if isinstance(doc, dict) and doc.get("file"):
                        files.append(doc["file"])
            result.tried.append(SiteAttempt(host, "done", "", rungs=[x for rr in r.results for x in rr.rungs if x not in ("primary",)][:12], repaired=len(r.repaired)))
            return files, pb, agent
        except BA.StepFailed as e:
            limit = "free limit" in e.why or "limit" in e.why.lower() and "credit" in e.why.lower()
            result.tried.append(SiteAttempt(host, "limit" if limit else "failed", e.why[:160], e.rungs[:10]))
            try:
                agent.backend.end()
            except Exception:  # noqa: BLE001
                pass
            return None
        except BA.NeedOwner as e:
            result.tried.append(SiteAttempt(host, "needs_owner", f"{e.kind}: {e.detail}"[:160]))
            if e.kind == "takeover":                                       # D79: the owner is needed at the screen and the task is paused for him: the next site is NOT tried,
                result.needs_owner = {"kind": "takeover", "site": host, "detail": e.detail[:160], "takeover": e.state.get("takeover")}      # and the browser session stays (his page)
                raise
            try:
                agent.backend.end()
            except Exception:  # noqa: BLE001
                pass
            if e.kind in ("grant",):
                raise
            return None
        except BA.NeedApproval as e:
            result.tried.append(SiteAttempt(host, "needs_owner", "held for your PIN: " + str(e)[:100]))
            result.needs_owner = {"kind": "approval", "site": host, "detail": str(e)[:160]}
            raise

    # ---- everything ---------------------------------------------------------------------------------------------------
    def run(self, topic: str, slides: int = 3, theme: str = "", formats=("pptx", "pdf"), edits: list = None, only=None, use_outline: bool = True) -> DeckResult:
        """Make the deck. Returns a DeckResult; its `facts` hold only what was verified. Never raises for a site that failed: it tries the next one."""
        topic = " ".join(str(topic or "").split())[:400]
        if len(topic) < 3:
            raise FlowError("what should the deck be about?")
        slides = max(1, min(int(slides or 3), 20))
        result = DeckResult(False)
        outline, provider = "", ""
        if use_outline and self.outline_fn:
            got = self.outline_fn(topic, slides)
            if got:
                outline, provider = got
                outline = re.sub(r"\s*\n+\s*", " | ", outline).strip(" |")[:OUTLINE_MAX]        # a job's text is one line of at most 500 characters
        rows = ranked_sites(self.shortlist, only)
        if not rows:
            result.summary = "No site on the shortlist could be used."
            return result
        for row in rows:
            if self.stop():
                raise BA.Stopped("stopped")
            try:
                out = self.run_site(row, topic, slides, theme, formats, outline, edits or [], result)
            except BA.NeedOwner:
                result.summary = f"Waiting for you: {result.needs_owner['detail']}" if result.needs_owner else "Waiting for you."
                return result
            except BA.NeedApproval:
                result.summary = "A step on " + row["host"] + " is held for your PIN in the console."
                return result
            if out is None:
                continue
            paths, pb, agent = out
            verified, problems = [], []
            for path in paths:
                try:
                    info = verify_file(path)
                    verified.append((path, info))
                except FlowError as e:
                    problems.append(str(e))
            try:
                agent.backend.end()
            except Exception:  # noqa: BLE001
                pass
            if not verified:
                result.tried[-1].outcome, result.tried[-1].why = "failed", ("; ".join(problems) or "no file arrived")[:160]
                continue
            result.ok, result.site = True, row["host"]
            for path, info in verified:
                fact = {k: info[k] for k in ("name", "ext", "size", "sha256", "slides", "pages") if k in info}
                result.files.append(fact)
                result.local_paths.append(path)
                if "slides" in info:
                    ok_n = info["slides"] == slides
                    result.facts.append(f"{info['ext']} file of {info['size']} bytes with {info['slides']} slides" + ("" if ok_n else f" ({slides} were asked for)"))
                elif "pages" in info:
                    result.facts.append(f"pdf of {info['pages']} pages, {info['size']} bytes")
            result.facts.append(f"made on {row['host']} in the owner's Chrome" + (f"; outline written by {provider}" if provider else ""))
            first = result.local_paths[0]
            try:
                shown = promote(first, topic, self.deliverables_dir)
                result.shown = shown
                if self.show:
                    self.show(shown)
            except OSError:
                pass
            result.summary = "Done: " + "; ".join(result.facts)
            return result
        result.summary = "Every site on the list failed or hit its free limit."
        result.needs_owner = {"kind": "all_failed", "tried": [{"site": t.host, "outcome": t.outcome, "why": t.why} for t in result.tried]}
        self._ask_all_failed(result)
        return result

    def _ask_all_failed(self, result: DeckResult) -> None:
        """Only now, when every site has failed, the owner is asked, with a form that shows what blocked each one."""
        if not self.ask:
            return
        words = {"failed": "a step could not be done", "limit": "its free limit is used up", "needs_owner": "it needed you", "no_playbook": "JARVIS has no steps for it yet",
                 "no_grant": "waiting for your grant", "stopped": "stopped"}
        rows = [[t.host, f"{words.get(t.outcome, t.outcome)}: {t.why}"[:200]] for t in result.tried]
        spec = {"title": "JARVIS could not make the deck", "intro": "Each site on the list was tried. Nothing was bought or upgraded. What would you like to do?",
                "fields": [{"type": "info_table", "label": "What blocked each site", "rows": rows}],
                "buttons": [{"id": "later", "label": "Leave it", "style": "primary"}, {"id": "manual", "label": "I will do it myself", "style": "secondary"}]}
        try:
            self.ask(spec)
        except Exception:  # noqa: BLE001
            pass
