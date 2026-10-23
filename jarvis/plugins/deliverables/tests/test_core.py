"""Deliverables logic without LibreOffice or OfficeCLI: spec checks, theme choice, text fitting, deck build, HTML safety,
and the rule that only made files can be opened. The full pipeline is run by run_demo.py and live_check.py."""
import json
import zipfile
from pathlib import Path

import pytest

from jarvis.plugins.deliverables import animate, deck, report, spec, themes, tools

FIX = Path(__file__).parent / "fixtures"


def demo(name):
    return json.loads((FIX / f"demo_{name}.json").read_text(encoding="utf-8"))


def test_demo_specs_are_valid():
    d = spec.deck(demo("deck"))
    r = spec.report(demo("report"))
    assert len(d["slides"]) == 10 and len(d["sources"]) == 5
    assert len(r["sections"]) == 4


@pytest.mark.parametrize("bad,msg", [
    ({"title": "x", "slides": []}, "at least one slide"),
    ({"title": "x", "slides": [{"layout": "nope", "title": "a"}]}, "layout must be"),
    ({"title": "x", "slides": [{"title": "a", "bullets": ["see [S9]"]}]}, "not listed"),
    ({"title": "x", "sources": [{"id": "S1", "title": "t", "url": "javascript:alert(1)"}], "slides": [{"title": "a"}]}, "http"),
    ({"title": "x", "slides": [{"layout": "chart", "title": "c", "chart": {"labels": ["a"], "series": [{"values": [1, 2]}]}}]}, "values for"),
    ({"title": "x" * 500, "slides": [{"title": "a"}]}, "too long"),
])
def test_bad_specs_are_refused(bad, msg):
    with pytest.raises(spec.SpecError, match=msg):
        spec.deck(bad)


def test_theme_choice():
    assert themes.choose_theme("Cloud security for engineers").name == "midnight"
    assert themes.choose_theme("Annual business report").name == "paper"
    assert themes.choose_theme("Kids music workshop").name == "vivid"
    assert themes.choose_theme("something", "vivid").name == "vivid"
    assert themes.choose_theme("plain words").name == themes.DEFAULT_THEME


def test_fit_size_shrinks_long_text_and_reports_overflow():
    short, ok = deck.fit_size([("Short line", 0)], "Segoe UI", 10, 2, 24, 14)
    assert short == 24 and ok
    long_paras = [("word " * 80, 0)] * 6
    size, ok = deck.fit_size(long_paras, "Segoe UI", 10, 3, 24, 14)
    assert size == 14 and not ok


def test_text_is_still_measured_to_scale_when_the_font_file_is_missing(monkeypatch):
    monkeypatch.setattr(deck, "font_path", lambda family, bold=False: None)    # no theme font, no Segoe UI either
    w12, w24 = deck._width_pt("word " * 10, "Segoe UI", 12), deck._width_pt("word " * 10, "Segoe UI", 24)
    assert w12 > 50 and 1.8 < w24 / w12 < 2.2                              # was ~25 pt at every size
    size, ok = deck.fit_size([("word " * 80, 0)] * 6, "Segoe UI", 10, 3, 24, 14)
    assert size == 14 and not ok                                            # overflow is still reported


def test_deck_build_without_external_tools(tmp_path):
    s = spec.deck(demo("deck"))
    out = tmp_path / "d.pptx"
    info = deck.build(s, str(out))
    assert info["slides"] == 11                               # 10 + one Sources slide before the closing slide
    assert all(f["fits"] for f in info["fit"])
    with zipfile.ZipFile(out) as z:
        charts = [n for n in z.namelist() if n.startswith("ppt/charts/chart")]
        xml = z.read(charts[0]).decode()
    import re
    ids = [int(v) for v in re.findall(r'<c:(?:axId|crossAx) val="(-?\d+)"', xml)]
    assert ids and all(0 <= v <= 0x7FFFFFFF for v in ids)     # schema-valid axis ids
    assert not any(a["kind"] == "table" for a in info["animations"])


def test_animation_commands_shape():
    cmds = animate.commands(3, "push", [{"slide": 2, "kind": "shape", "id": 5, "effect": "fade"},
                                        {"slide": 2, "kind": "chart", "id": 7, "effect": "wipe"}], {1})
    assert cmds[0]["props"]["transition"] == "fade" and cmds[1]["props"]["transition"] == "push"
    adds = [c for c in cmds if c["command"] == "add"]
    assert adds[0]["parent"] == "/slide[2]/shape[@id=5]" and adds[1]["parent"] == "/slide[2]/chart[@id=7]"
    assert adds[0]["props"]["delay"] == "250" and adds[1]["props"]["delay"] == "120"


def test_report_html_escapes_untrusted_text():
    s = spec.report({"title": "<script>alert(1)</script>", "sources": [], "sections": [
        {"heading": "H <img src=x onerror=alert(1)>", "paragraphs": ["[click](javascript:alert(1)) and <b>raw</b>"]}]})
    html, _ = report.render_html(s)
    assert "<script>alert(1)</script>" not in html
    assert "<img src=x" not in html
    assert 'href="javascript' not in html                    # left as plain text, never a link
    assert "&lt;b&gt;raw&lt;/b&gt;" in html


def test_report_citations_link_to_references():
    html, _ = report.render_html(spec.report(demo("report")))
    assert 'href="#src-S1"' in html and 'id="src-S1"' in html


def test_only_made_files_can_be_opened(tmp_path, monkeypatch):
    monkeypatch.setattr(tools, "OUT_ROOT", tmp_path)
    made = tmp_path / "job" / "a.pdf"
    made.parent.mkdir()
    made.write_bytes(b"%PDF")
    assert tools.inside_outputs(str(made)) == made.resolve()
    outside = tmp_path.parent / "elsewhere.pdf"
    outside.write_bytes(b"%PDF")
    with pytest.raises(tools.ToolError):
        tools.inside_outputs(str(outside))
    with pytest.raises(tools.ToolError):
        tools.inside_outputs(str(tmp_path / "job" / "missing.pdf"))
    exe = tmp_path / "job" / "x.exe"
    exe.write_bytes(b"MZ")
    with pytest.raises(tools.ToolError):
        tools.inside_outputs(str(exe))


def test_new_job_dir_never_reuses(tmp_path, monkeypatch):
    monkeypatch.setattr(tools, "OUT_ROOT", tmp_path)
    a, b = tools.new_job_dir("Same Title"), tools.new_job_dir("Same Title")
    assert a != b and a.is_dir() and b.is_dir()


def test_a_stats_slide_with_more_than_four_items_is_refused_not_cut():
    s = {"title": "x", "slides": [{"layout": "stats", "title": "s", "items": [{"value": str(i), "label": f"l{i}"} for i in range(6)]}]}
    with pytest.raises(spec.SpecError, match="at most 4"):
        spec.deck(s)
    four = {"title": "x", "slides": [{"layout": "stats", "title": "s", "items": [{"value": str(i), "label": f"l{i}"} for i in range(4)]}]}
    assert len(spec.deck(four)["slides"][0]["items"]) == 4


def test_missing_edge_is_reported_plainly(monkeypatch):
    import playwright.sync_api as psa

    class NoEdge:
        def __enter__(self):
            class C:
                def launch(self, **kw):
                    raise RuntimeError("Chromium distribution 'msedge' is not found. Run \"npx playwright install msedge\"")
            return type("P", (), {"chromium": C()})()

        def __exit__(self, *e):
            return False
    monkeypatch.setattr(psa, "sync_playwright", NoEdge)
    with pytest.raises(tools.ToolError, match="Microsoft Edge could not be started") as e:
        report.print_pdf("<p>x</p>", "x.pdf", "t")
    assert "npx" not in str(e.value)


@pytest.mark.parametrize("fn,bad,msg", [
    (spec.deck, {"title": "T", "slides": [{"layout": 5, "title": "x"}]}, r"slides\[0\]\.layout must be text"),
    (spec.deck, {"title": "T", "slides": [{"title": "x", "bullets": [{"text": "a", "level": "high"}]}]}, r"level must be a whole number"),
    (spec.deck, {"title": "T", "slides": [{"title": "x", "bullets": [{"text": "a", "level": [1]}]}]}, r"level must be a whole number"),
    (spec.deck, {"title": "T", "slides": [{"layout": "chart", "title": "x", "chart": {"type": 3, "labels": ["a"], "series": [{"values": [1]}]}}]}, r"chart\.type must be text"),
    (spec.deck, {"title": "T", "slides": [{"layout": "chart", "title": "x", "chart": {"labels": [], "series": [5]}}]}, r"series\[0\] must be an object"),
    (spec.report, {"title": "T", "sections": [{"heading": "h", "kind": 3}]}, r"kind must be text"),
    (spec.report, {"title": "T", "sections": [{"heading": "h", "bullets": [{"text": "a", "level": "x"}]}]}, r"level must be a whole number"),
])
def test_wrongly_typed_fields_are_refused_with_spec_error(fn, bad, msg):
    """Bench I-019: these raised AttributeError / ValueError, which the plugin entry (it catches SpecError) did not report."""
    with pytest.raises(spec.SpecError, match=msg):
        fn(bad)


@pytest.mark.parametrize("v", ["nan", "inf", "-inf", 1e309, float("nan")])
def test_non_finite_chart_values_are_refused(v):
    """Bench I-020: float() accepts these; they passed the spec and failed later inside the chart writer."""
    s = {"title": "T", "slides": [{"layout": "chart", "title": "c", "chart": {"labels": ["a", "b"], "series": [{"values": [v, 1]}]}}]}
    with pytest.raises(spec.SpecError, match="finite"):
        spec.deck(s)


def test_control_characters_never_show_as_codes_on_slides(tmp_path):
    """Bench I-021: python-pptx wrote vertical tab, form feed, NUL, bell and SOH as visible _x000B_-style codes."""
    import re

    from pptx import Presentation
    s = {"title": "T", "slides": [{"layout": "bullets", "title": "Extracted\x0btext", "notes": "speaker\x01notes",
                                   "bullets": ["line one\x0cline two", "nul\x00here", "bell\x07sound\x85", "tab\tkept"]}]}
    checked = spec.deck(s)
    assert checked["slides"][0]["title"] == "Extracted text" and checked["slides"][0]["bullets"][3]["text"] == "tab\tkept"
    assert spec.report({"title": "T", "sections": [{"heading": "h", "paragraphs": ["a\r\nb\x1bc"]}]})["sections"][0]["paragraphs"] == ["a\nb c"]
    out = tmp_path / "c.pptx"
    deck.build(checked, str(out))
    prs = Presentation(str(out))
    texts = [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]
    texts += [sl.notes_slide.notes_text_frame.text for sl in prs.slides if sl.has_notes_slide]
    assert not [m for x in texts for m in re.findall(r"_x00[0-9A-Fa-f]{2}_", x)]
    for w in ("Extracted text", "line two", "nul here", "bell sound", "speaker notes"):
        assert any(w in x for x in texts), w


def test_report_html_turns_images_into_links_and_loads_nothing_remote():
    """Bench I-022: a markdown image became <img src="http://..."> in the report.html saved next to the PDF."""
    import re
    s = spec.report({"title": "T", "sections": [{"heading": "h",
                     "paragraphs": ["![pixel](http://127.0.0.1:9/track.png?u=owner) and ![x](javascript:alert(1)) and ![d](data:image/png;base64,AAAA)"],
                     "callout": "![beacon](https://tracker.example/p.gif)"}]})
    html, _ = report.render_html(s)
    assert not re.findall(r"<(?:img|iframe|script|link|object|embed)\b", html)
    assert 'href="http://127.0.0.1:9/track.png?u=owner">pixel</a>' in html
    assert 'href="https://tracker.example/p.gif">beacon</a>' in html
    assert all(h.startswith(("https://", "http://", "#")) for h in re.findall(r'href="([^"]*)"', html))


@pytest.mark.parametrize("damage", ["cut", "junk"])
def test_qa_on_a_damaged_pdf_says_not_passed(tmp_path, damage):
    """Bench I-029: pypdf raised a raw PdfStreamError on a cut PDF or junk behind a %PDF header."""
    from pypdf import PdfWriter

    from jarvis.plugins.deliverables import qa
    good = tmp_path / "g.pdf"
    w = PdfWriter()
    w.add_blank_page(200, 200)
    with open(good, "wb") as f:
        w.write(f)
    data = good.read_bytes()
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(data[: len(data) // 2] if damage == "cut" else b"%PDF-1.7\n" + bytes((i * 91 + 7) % 256 for i in range(4000)))
    res = qa.check_report(bad, spec.report(demo("report")), tmp_path)
    assert res["passed"] is False and res["warnings"]


def test_new_job_dir_is_safe_when_called_at_the_same_moment(tmp_path, monkeypatch):
    """Bench I-032: exists() then mkdir() raced; 16 of 40 simultaneous calls raised FileExistsError."""
    import threading
    monkeypatch.setattr(tools, "OUT_ROOT", tmp_path)
    real_mkdir, gate, seen = Path.mkdir, threading.Barrier(8), threading.local()

    def slow_mkdir(self, *a, **kw):              # every thread passes its existence check before any folder is made
        if self.parent == tmp_path and not getattr(seen, "once", False):
            seen.once = True
            gate.wait(timeout=10)
        return real_mkdir(self, *a, **kw)
    monkeypatch.setattr(Path, "mkdir", slow_mkdir)
    dirs, errs = [], []

    def mk():
        try:
            dirs.append(tools.new_job_dir("Same title"))
        except Exception as e:  # noqa: BLE001
            errs.append(type(e).__name__)
    ts = [threading.Thread(target=mk) for _ in range(8)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert errs == [] and len(set(dirs)) == 8 and all(d.is_dir() for d in dirs)
