"""Safe document reading and writing for the Work and Learn modes.

Reading is limited to folders the owner names (default: Documents, Downloads, Desktop, the JARVIS outputs folder, the study
vault), refuses anything that looks like a credential, and never executes what it reads. Writing goes only to one output
folder and never overwrites. Everything returned is plain text for a model to read; it is untrusted data, not instructions.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

MAX_READ_BYTES = 25 * 1024 * 1024
DEFAULT_CHARS = 20_000
MAX_CHARS = 100_000
MAX_PDF_PAGES = 200
MAX_WRITE_CHARS = 200_000
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml", ".yml", ".toml", ".ini", ".rst", ".xml"}
HTML_SUFFIXES = {".html", ".htm"}
BLOCKED_DIR_NAMES = {".ssh", ".gnupg", ".aws", ".azure", "appdata", "$recycle.bin", "system volume information",
                     "windows", "program files", "program files (x86)", "jarvis-secrets"}
BLOCKED_NAME = re.compile(r"(^\.env|secret|token|password|passwd|credential|id_rsa|id_ed25519|\.pem$|\.key$|\.pfx$|\.kdbx$|wallet)", re.I)
NAME_OK = re.compile(r"[^A-Za-z0-9 _.\-]+")
FORMATS = ("md", "txt", "docx")


class DocError(Exception):
    """A refusal or failure whose message is safe to show to the user."""


def read_roots() -> list[Path]:
    env = os.environ.get("DOC_TOOLS_READ_ROOTS", "").strip()
    if env:
        raw = [p for p in env.split(os.pathsep) if p.strip()]
    else:
        home = Path(os.environ.get("USERPROFILE") or Path.home())
        raw = [str(home / "Documents"), str(home / "Downloads"), str(home / "Desktop"), r"D:\jarvis\outputs",
               os.environ.get("LEARN_VAULT_DIR", r"D:\jarvis\learn-vault")]
    out: list[Path] = []
    for p in raw:
        try:
            out.append(Path(p.strip().strip('"')).resolve())
        except OSError:
            continue
    return out


def out_dir() -> Path:
    return Path(os.environ.get("DOC_TOOLS_OUT_DIR", r"D:\jarvis\outputs\documents")).resolve()


def resolve_readable(path_str: str) -> Path:
    raw = str(path_str or "").strip().strip('"')
    if not raw:
        raise DocError("No file path was given.")
    path = Path(raw)
    if not path.is_absolute():
        raise DocError("Give the full path of the file.")
    try:
        real = path.resolve(strict=True)
    except OSError:
        raise DocError("That file does not exist.")
    if not real.is_file():
        raise DocError("That path is not a file.")
    parts = {p.lower() for p in real.parts}
    if parts & BLOCKED_DIR_NAMES or BLOCKED_NAME.search(real.name):
        raise DocError("That file is protected and will not be read.")
    if not any(real == r or r in real.parents for r in read_roots()):
        raise DocError("That file is outside the folders this tool may read (Documents, Downloads, Desktop, outputs, the study vault).")
    if real.stat().st_size > MAX_READ_BYTES:
        raise DocError("That file is too large to read (25 MB limit).")
    return real


NO_TEXT_LAYER = "no text layer: the PDF may be scanned; OCR is not available here"


def _read_pdf(path: Path) -> tuple[str, int, int]:
    """(text, pages read, pages that had no extractable text). Scanned or image-only pages give no text: there is no OCR
    here, so the caller says so instead of returning silent empty text."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocError("That PDF is password protected.")
    pages, read, empty = [], 0, 0
    for i, page in enumerate(reader.pages):
        if i >= MAX_PDF_PAGES:
            pages.append(f"[stopped after {MAX_PDF_PAGES} pages]")
            break
        text = page.extract_text() or ""
        read += 1
        empty += not text.strip()
        pages.append(text)
    return "\n\n".join(pages), read, empty


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def read_document(path_str: str, max_chars: int = DEFAULT_CHARS) -> dict[str, Any]:
    real = resolve_readable(path_str)
    limit = max(500, min(int(max_chars or DEFAULT_CHARS), MAX_CHARS))
    suffix = real.suffix.lower()
    note = ""
    try:
        if suffix == ".pdf":
            text, read, empty = _read_pdf(real)
            kind = "pdf"
            if read and empty == read:
                note = NO_TEXT_LAYER
            elif empty:
                note = f"{empty} of {read} pages have no text layer (they may be scanned; OCR is not available here)"
        elif suffix == ".docx":
            text, kind = _read_docx(real), "docx"
        elif suffix in HTML_SUFFIXES:
            from bs4 import BeautifulSoup

            html = real.read_text(encoding="utf-8", errors="replace")
            text, kind = BeautifulSoup(html, "html.parser").get_text("\n"), "html"
        elif suffix in TEXT_SUFFIXES or suffix == "":
            text, kind = real.read_text(encoding="utf-8", errors="replace"), "text"
        else:
            raise DocError(f"This tool reads text, Markdown, CSV, JSON, HTML, PDF and Word files, not '{suffix or 'no extension'}'.")
    except DocError:
        raise
    except Exception as exc:                                   # a broken PDF or docx: say so without echoing internals
        raise DocError(f"That file could not be read ({type(exc).__name__}).")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    out = {"name": real.name, "kind": kind, "chars_total": len(text), "truncated": len(text) > limit, "text": text[:limit]}
    if note:
        out["note"] = note
    return out


def safe_name(name: str) -> str:
    base = NAME_OK.sub("", str(name or "")).strip(" .-_")
    base = re.sub(r"\.(md|txt|docx)$", "", base, flags=re.I)[:60].strip(" .-_")
    if not base:
        raise DocError("Give the document a name using letters, numbers, spaces, dots, dashes or underscores.")
    return base


def _unique(directory: Path, base: str, ext: str) -> Path:
    candidate = directory / f"{base}.{ext}"
    n = 2
    while candidate.exists():
        candidate = directory / f"{base}-{n}.{ext}"
        n += 1
        if n > 999:
            raise DocError("Too many documents with that name.")
    return candidate


def _markdown_to_docx(content: str, target: Path) -> None:
    import docx

    document = docx.Document()
    for raw in content.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            document.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", line):
            document.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            document.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            document.add_paragraph(line.strip())
    document.save(str(target))


def write_document(name: str, content: str, fmt: str = "md") -> dict[str, Any]:
    fmt = str(fmt or "md").lower().lstrip(".")
    if fmt not in FORMATS:
        raise DocError(f"Format must be one of: {', '.join(FORMATS)}.")
    text = str(content or "")
    if not text.strip():
        raise DocError("There is no content to write.")
    if len(text) > MAX_WRITE_CHARS:
        raise DocError(f"The content is too long ({len(text)} characters; at most {MAX_WRITE_CHARS}).")
    base = safe_name(name)
    directory = out_dir()
    directory.mkdir(parents=True, exist_ok=True)
    target = _unique(directory, base, fmt)
    if fmt == "docx":
        _markdown_to_docx(text, target)
    else:
        target.write_text(text if text.endswith("\n") else text + "\n", encoding="utf-8", newline="")
    return {"path": str(target), "name": target.name, "format": fmt, "bytes": target.stat().st_size}


def list_outputs(limit: int = 50) -> list[dict[str, Any]]:
    directory = out_dir()
    if not directory.is_dir():
        return []
    files = sorted((p for p in directory.iterdir() if p.is_file()), key=lambda p: p.stat().st_mtime, reverse=True)
    return [{"name": p.name, "path": str(p), "bytes": p.stat().st_size} for p in files[: max(1, min(int(limit), 200))]]
