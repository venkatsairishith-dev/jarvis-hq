"""Check a finished deck or report the way a careful person would, and save page images for a visual look.

Every check reports what it found; nothing is silently fixed. A deliverable "passes" only when all hard checks pass.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pdfplumber
import pypdfium2 as pdfium
from PIL import Image
from pypdf import PdfReader

from . import animate, tools


def page_images(pdf: Path, out_dir: Path, prefix: str, scale: float = 0.6, limit: int = 60) -> list[str]:
    out_dir.mkdir(exist_ok=True)
    doc = pdfium.PdfDocument(str(pdf))
    paths = []
    try:
        for i in range(min(len(doc), limit)):
            img = doc[i].render(scale=scale).to_pil()
            p = out_dir / f"{prefix}-{i + 1:02d}.png"
            img.save(p)
            paths.append(str(p))
    finally:
        doc.close()
    return paths


def contact_sheet(images: list[str], path: Path, cols: int = 4, width: int = 360) -> str | None:
    if not images:
        return None
    thumbs = []
    for p in images:
        im = Image.open(p)
        thumbs.append(im.resize((width, int(im.height * width / im.width))))
    h = max(t.height for t in thumbs)
    rows = (len(thumbs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (width + 12) + 12, rows * (h + 12) + 12), "white")
    for i, t in enumerate(thumbs):
        sheet.paste(t, (12 + (i % cols) * (width + 12), 12 + (i // cols) * (h + 12)))
    sheet.save(path)
    return str(path)


def words_off_page(pdf: Path, tolerance: float = 1.0) -> list[dict]:
    """Text drawn outside the page area: the clearest sign of overflow in a rendered page."""
    bad = []
    with pdfplumber.open(str(pdf)) as doc:
        for n, page in enumerate(doc.pages, 1):
            for w in page.extract_words():
                if w["x0"] < -tolerance or w["x1"] > page.width + tolerance or w["top"] < -tolerance or w["bottom"] > page.height + tolerance:
                    bad.append({"page": n, "text": w["text"][:30]})
    return bad[:20]


def check_deck(deck: Path, build_info: dict, spec: dict, work: Path) -> dict:
    from pptx import Presentation
    checks, warnings = {}, []
    prs = Presentation(str(deck))                       # reopens cleanly
    checks["reopens"] = True
    checks["slide_count"] = len(prs.slides) == build_info["slides"]
    notes_wanted = sum(1 for s in spec["slides"] if s["notes"])
    notes_found = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    checks["speaker_notes"] = notes_found >= notes_wanted
    overflow = [f for f in build_info["fit"] if not f["fits"]]
    checks["text_fits"] = not overflow
    if overflow:
        warnings += [f"slide {f['slide']}: '{f['box']}' text is too long even at {f['size']:.0f} pt" for f in overflow]
    with zipfile.ZipFile(deck) as z:
        xml = [z.read(n).decode("utf-8", "replace") for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    trans = sum("<p:transition" in x for x in xml)
    timing = sum("<p:timing" in x for x in xml)
    checks["transitions"] = (trans == len(xml)) if spec["animate"] else True
    checks["animations"] = (timing >= len({a["slide"] for a in build_info["animations"]})) if spec["animate"] else True
    v = animate.validate(deck)
    checks["schema_valid"] = v["valid"]
    if not v["valid"]:
        warnings += v["errors"]
    rendered = None
    try:
        rendered = tools.soffice_to_pdf(deck, work)
        checks["renders"] = True
        with pdfplumber.open(str(rendered)) as doc:
            checks["rendered_pages"] = len(doc.pages) == build_info["slides"]
        off = words_off_page(rendered)
        checks["nothing_off_page"] = not off
        if off:
            warnings.append(f"text outside the slide: {off[:3]}")
        imgs = page_images(rendered, work / "previews", "slide")
        sheet = contact_sheet(imgs, work / "previews" / "contact-sheet.png")
    except tools.ToolError as e:
        checks["renders"] = False
        warnings.append(str(e))
        imgs, sheet = [], None
    return {"passed": all(checks.values()), "checks": checks, "warnings": warnings,
            "counts": {"slides": build_info["slides"], "transitions": trans, "slides_with_animation": timing,
                       "notes": notes_found},
            "previews": imgs, "contact_sheet": sheet, "rendered_pdf": str(rendered) if rendered else None,
            "renderer": "LibreOffice 26.2.6 (not PowerPoint: effects may look simpler there)"}


def check_report(pdf: Path, spec: dict, work: Path) -> dict:
    checks, warnings = {}, []
    reader = PdfReader(str(pdf))
    pages = len(reader.pages)
    checks["opens"] = pages > 0
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    flat = re.sub(r"\s+", " ", text)
    missing = [s["heading"] for s in spec["sections"] if re.sub(r"\s+", " ", s["heading"])[:40] not in flat]
    checks["searchable_text_and_headings"] = bool(text.strip()) and not missing
    if missing:
        warnings.append(f"headings not found in the text layer: {missing[:3]}")
    checks["title_present"] = spec["title"][:40] in flat
    try:
        outline = reader.outline
        n_outline = len(outline) if outline else 0
    except Exception:  # noqa: BLE001
        n_outline = 0
    checks["bookmarks"] = n_outline > 0
    uris, internal = [], 0
    for p in reader.pages:
        for a in p.get("/Annots") or []:
            a = a.get_object()
            if a.get("/Subtype") != "/Link":
                continue
            act = a.get("/A")
            if act and act.get_object().get("/URI"):
                uris.append(str(act.get_object()["/URI"]))
            else:
                internal += 1
    want_urls = {s["url"] for s in spec["sources"] if s["url"]}
    checks["source_links"] = want_urls <= set(uris)
    if not checks["source_links"]:
        warnings.append(f"missing links: {sorted(want_urls - set(uris))[:3]}")
    fonts, unembedded = set(), set()
    for p in reader.pages:
        res = p.get("/Resources")
        fdict = res.get_object().get("/Font") if res else None
        for f in (fdict.get_object().values() if fdict else []):
            f = f.get_object()
            name = str(f.get("/BaseFont", "?"))
            fonts.add(name)
            desc = f.get("/FontDescriptor")
            if desc is None and f.get("/DescendantFonts"):
                desc = f["/DescendantFonts"][0].get_object().get("/FontDescriptor")
            if desc is not None and f.get("/Subtype") != "/Type3":
                d = desc.get_object()
                if not any(k in d for k in ("/FontFile", "/FontFile2", "/FontFile3")):
                    unembedded.add(name)
    checks["fonts_embedded"] = not unembedded
    if unembedded:
        warnings.append(f"fonts not embedded: {sorted(unembedded)}")
    root = reader.trailer["/Root"].get_object()
    checks["tagged"] = "/StructTreeRoot" in root
    off = words_off_page(pdf)
    checks["nothing_off_page"] = not off
    imgs = page_images(pdf, work / "previews", "page", scale=0.8)
    sheet = contact_sheet(imgs, work / "previews" / "contact-sheet.png", cols=5, width=300)
    return {"passed": all(checks.values()), "checks": checks, "warnings": warnings,
            "counts": {"pages": pages, "links": len(uris), "internal_links": internal, "bookmarks": n_outline, "fonts": len(fonts)},
            "previews": imgs, "contact_sheet": sheet}
