"""Build a PowerPoint deck from a checked spec with python-pptx (native text, charts, tables and notes).

Text is measured with the real font files before saving: each box gets the largest size (down to a floor) at which its
text fits, and anything that still does not fit is reported, never hidden. Transitions and entrance animations are added
afterwards by animate.py (OfficeCLI), using the shape ids recorded here.
"""
from __future__ import annotations

import datetime as _dt
import re
from functools import lru_cache

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .spec import CITE
from .themes import Theme, choose_theme, font_path

W, H = 13.333, 7.5                 # 16:9, inches
MX = 0.75                          # side margin
INSET = 0.08                       # text box inset, inches each side
_CHARTS = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
           "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


# ---------- text measuring ----------

@lru_cache(maxsize=64)
def _font(family: str, bold: bool, size10: int):
    path = font_path(family, bold) or font_path("Segoe UI", bold)
    return ImageFont.truetype(path, size10) if path else ImageFont.load_default()


def _width_pt(text: str, family: str, size: float, bold: bool = False) -> float:
    return _font(family, bold, int(size * 10)).getlength(text) / 10.0


def _lines(text: str, family: str, size: float, width_pt: float, bold: bool = False) -> int:
    n = 0
    for para in (text or "").split("\n"):
        words, cur = para.split(" "), ""
        n += 1
        for w in words:
            trial = (cur + " " + w).strip()
            if _width_pt(trial, family, size, bold) <= width_pt or not cur:
                cur = trial
            else:
                n += 1
                cur = w
    return n


def fit_size(paras: list[tuple[str, int]], family: str, width_in: float, height_in: float, start: float, floor: float,
             bold: bool = False, spacing: float = 1.2, gap: float = 0.45, indent_in: float = 0.32) -> tuple[float, bool]:
    """Largest font size (pt) from start down to floor at which all paragraphs fit the box; (size, fits)."""
    size = start
    while True:
        need = 0.0
        for text, level in paras:
            w = (width_in - 2 * INSET - (indent_in * (level + 1) if level >= 0 else 0)) * 72
            need += _lines(text, family, size, w, bold) * size * spacing + size * gap
        need -= size * gap if paras else 0
        if need <= (height_in - 2 * INSET) * 72 or size <= floor:
            return size, need <= (height_in - 2 * INSET) * 72
        size -= 1


# ---------- low-level shape helpers ----------

class Builder:
    def __init__(self, spec: dict):
        self.spec = spec
        self.theme: Theme = choose_theme(" ".join([spec["title"], spec.get("subtitle", ""), spec.get("topic", "")]), spec.get("theme"))
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.src_num = {s["id"]: i + 1 for i, s in enumerate(spec["sources"])}
        self.fit_report: list[dict] = []
        self.anim: list[dict] = []          # {"slide": n, "kind": "shape"|"chart", "id": shape id, "effect": ...}
        self.slide_no = 0

    # backgrounds and decoration
    def _bg(self, slide, color: str | None = None):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color or self.theme.bg)

    def _rect(self, slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, alpha: int | None = None, line=False):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        if alpha is not None:          # a:alpha inside the solid fill colour (percent * 1000)
            clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
        if not line:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _text(self, slide, x, y, w, h, paras, *, family=None, size=20, floor=12, color=None, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, bullets=False, italic=False, name="text", spacing=1.2, gap=0.45):
        """paras: list of (text, level). Level -1 means no bullet. Citation markers become superscript numbers."""
        family = family or self.theme.body_font
        plain = [(CITE.sub(lambda m: f"[{self.src_num.get(m.group(1), '?')}]", t), (lvl if bullets else -1)) for t, lvl in paras]
        size, fits = fit_size(plain, family, w, h, size, floor, bold, spacing, gap)
        self.fit_report.append({"slide": self.slide_no, "box": name, "size": size, "fits": fits})
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.vertical_anchor = anchor
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(INSET))
        for i, (text, level) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            if i:
                p.space_before = Pt(size * gap)
            if bullets:
                self._bullet(p, level, size)
            self._runs(p, text, family, size, color or self.theme.text, bold, italic)
        return box

    def _bullet(self, p, level: int, size: float):
        pPr = p._p.get_or_add_pPr()
        ind = Inches(0.32) * (level + 1)
        pPr.set("marL", str(int(ind)))
        pPr.set("indent", str(-int(Inches(0.3))))
        for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar"):
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", self.theme.accent if level == 0 else self.theme.muted)
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if level == 0 else "–")

    def _runs(self, p, text, family, size, color, bold, italic):
        pos = 0
        for m in list(CITE.finditer(text)) + [None]:
            end = m.start() if m else len(text)
            if end > pos:
                r = p.add_run()
                r.text = text[pos:end]
                self._style(r, family, size, color, bold, italic)
            if m:
                r = p.add_run()
                r.text = f"[{self.src_num.get(m.group(1), '?')}]"
                # full size + superscript: renderers already shrink superscript (~60%); shrinking it here too made the
                # numbers unreadable dots (seen in the LibreOffice render)
                self._style(r, family, max(12, size), self.theme.accent, True, False)
                r.font._rPr.set("baseline", "30000")
                pos = m.end()

    @staticmethod
    def _style(run, family, size, color, bold, italic):
        f = run.font
        f.name, f.size, f.bold, f.italic = family, Pt(size), bold, italic
        f.color.rgb = rgb(color)

    # ---------- slide scaffolding ----------

    def _new(self, layout: str):
        slide = self.prs.slides.add_slide(self.blank)
        self.slide_no += 1
        self._bg(slide)
        return slide

    def _title(self, slide, title, subtitle=""):
        self._rect(slide, MX, 0.62, 0.09, 0.62, self.theme.accent)
        self._text(slide, MX + 0.25, 0.45, W - 2 * MX - 0.25, 0.95, [(title, -1)], family=self.theme.heading_font,
                   size=34, floor=22, bold=self.theme.heading_font != "Segoe UI Semibold", anchor=MSO_ANCHOR.MIDDLE, name="title")
        if subtitle:
            self._text(slide, MX + 0.25, 1.35, W - 2 * MX - 0.25, 0.5, [(subtitle, -1)], size=16, floor=12,
                       color=self.theme.muted, name="subtitle")

    def _footer(self, slide, cited: list[str]):
        t = self.theme
        self._text(slide, MX, 6.98, 6.5, 0.36, [(self.spec["title"], -1)], size=10, floor=8, color=t.muted, name="footer")
        self._text(slide, W - MX - 1.6, 6.98, 1.6, 0.36, [(f"{self.slide_no}", -1)], size=10, floor=8, color=t.muted,
                   align=PP_ALIGN.RIGHT, name="page")
        if cited:
            line = "  ".join(f"[{self.src_num[s]}] {self._src_short(s)}" for s in cited)
            self._text(slide, MX, 6.55, W - 2 * MX, 0.42, [(line, -1)], size=10, floor=7, color=t.muted, name="footnote")

    def _src_short(self, sid: str) -> str:
        s = next(x for x in self.spec["sources"] if x["id"] == sid)
        host = re.sub(r"^https?://(www\.)?", "", s["url"]).split("/")[0] if s["url"] else ""
        bits = [s["title"][:70] + ("…" if len(s["title"]) > 70 else ""), host, s["date"]]
        return ", ".join(b for b in bits if b)

    def _cited(self, sl: dict) -> list[str]:
        found = []
        texts = [sl["title"], sl["subtitle"], sl["quote"]] + [b["text"] for b in sl["bullets"] + sl["left"] + sl["right"]] \
            + [it["text"] for it in sl["items"]]
        for t in texts:
            for m in CITE.findall(t or ""):
                if m not in found:
                    found.append(m)
        for s in sl["sources"]:
            if s not in found:
                found.append(s)
        return found[:4]

    def _animate(self, shape, effect="fade", kind="shape"):
        self.anim.append({"slide": self.slide_no, "kind": kind, "id": shape.shape_id, "effect": effect})

    def _notes(self, slide, sl: dict, cited: list[str]):
        text = sl.get("notes") or ""
        if cited:
            text = (text + "\n\n" if text else "") + "Sources: " + "; ".join(
                f"[{self.src_num[s]}] {next(x for x in self.spec['sources'] if x['id'] == s)['url'] or self._src_short(s)}" for s in cited)
        if text:
            slide.notes_slide.notes_text_frame.text = text

    # ---------- layouts ----------

    def title_slide(self, sl: dict):
        t, s = self.theme, self._new("title")
        self._bg(s, t.bg)
        self._rect(s, 8.9, -1.6, 6.2, 6.2, t.accent, MSO_SHAPE.OVAL, alpha=18)
        self._rect(s, 10.4, 3.7, 4.6, 4.6, t.accent2, MSO_SHAPE.OVAL, alpha=16)
        self._rect(s, MX, 2.35, 1.1, 0.1, t.accent)
        tb = self._text(s, MX, 2.6, 8.6, 2.0, [(sl["title"] or self.spec["title"], -1)], family=t.heading_font, size=46,
                        floor=28, bold=t.heading_font != "Segoe UI Semibold", anchor=MSO_ANCHOR.TOP, name="title", spacing=1.05)
        sub = sl["subtitle"] or self.spec["subtitle"]
        if sub:
            st = self._text(s, MX, 4.7, 8.6, 1.0, [(sub, -1)], size=22, floor=14, color=t.muted, name="subtitle")
            self._animate(st, "fade")
        meta = "  ·  ".join(x for x in (self.spec["author"], _dt.date.today().strftime("%d %B %Y")) if x)
        self._text(s, MX, 6.3, 8.0, 0.45, [(meta, -1)], size=13, floor=10, color=t.muted, name="meta")
        self._notes(s, sl, [])
        return tb

    def section(self, sl: dict, index: int):
        t, s = self.theme, self._new("section")
        self._rect(s, 0, 0, 4.2, H, t.surface)
        num = self._text(s, 0.6, 2.4, 3.2, 2.0, [(f"{index:02d}", -1)], family=t.heading_font, size=96, floor=60,
                         color=t.accent, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="number")
        self._text(s, 4.9, 2.5, W - 4.9 - MX, 1.5, [(sl["title"], -1)], family=t.heading_font, size=40, floor=26,
                   bold=t.heading_font != "Segoe UI Semibold", anchor=MSO_ANCHOR.BOTTOM, name="title")
        if sl["subtitle"]:
            self._text(s, 4.9, 4.1, W - 4.9 - MX, 1.2, [(sl["subtitle"], -1)], size=20, floor=14, color=t.muted, name="subtitle")
        self._animate(num, "zoom")
        self._notes(s, sl, [])

    def bullets(self, sl: dict):
        s = self._new("bullets")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        top = 1.95 if sl["subtitle"] else 1.6
        bottom = 6.45 if cited else 6.8
        body = self._text(s, MX + 0.25, top, W - 2 * MX - 0.25, bottom - top, [(b["text"], b["level"]) for b in sl["bullets"]],
                          size=24, floor=14, bullets=True, name="body")
        self._animate(body, "fade")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def two_column(self, sl: dict):
        t, s = self.theme, self._new("two_column")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        top, bottom = (1.95 if sl["subtitle"] else 1.6), (6.45 if cited else 6.8)
        cw = (W - 2 * MX - 0.5) / 2
        for k, (head, items) in enumerate(((sl["left_title"], sl["left"]), (sl["right_title"], sl["right"]))):
            x = MX + k * (cw + 0.5)
            self._rect(s, x, top, cw, bottom - top, t.surface, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
            y = top + 0.2
            if head:
                self._text(s, x + 0.25, y, cw - 0.5, 0.6, [(head, -1)], family=t.heading_font, size=22, floor=16,
                           color=t.accent if k == 0 else t.accent2, bold=True, name=f"col{k}_head")
                y += 0.7
            col = self._text(s, x + 0.25, y, cw - 0.5, bottom - y - 0.2, [(b["text"], b["level"]) for b in items],
                             size=20, floor=13, bullets=True, name=f"col{k}")
            self._animate(col, "fade")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def chart_slide(self, sl: dict):
        t, s = self.theme, self._new("chart")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        c = sl["chart"]
        top, bottom = (1.95 if sl["subtitle"] else 1.6), (6.45 if cited else 6.8)
        side = bool(sl["bullets"])
        cw = (W - 2 * MX) * (0.62 if side else 1.0)
        data = CategoryChartData()
        data.categories = c["labels"]
        for ser in c["series"]:
            data.add_series(ser["name"], ser["values"])
        gf = s.shapes.add_chart(_CHARTS[c["type"]], Inches(MX), Inches(top), Inches(cw), Inches(bottom - top - (0.4 if c["caption"] else 0)), data)
        self._style_chart(gf.chart, c)
        _fix_axis_ids(gf.chart)
        self._animate(gf, "wipe" if c["type"] != "pie" else "zoom", kind="chart")
        if c["caption"]:
            self._text(s, MX, bottom - 0.4, cw, 0.4, [(c["caption"], -1)], size=12, floor=9, color=t.muted, italic=True, name="caption")
        if side:
            body = self._text(s, MX + cw + 0.4, top, W - 2 * MX - cw - 0.4, bottom - top, [(b["text"], b["level"]) for b in sl["bullets"]],
                              size=20, floor=13, bullets=True, name="body")
            self._animate(body, "fade")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def _style_chart(self, chart, c: dict):
        t = self.theme
        chart.font.size = Pt(12)
        chart.font.name = t.body_font
        chart.font.color.rgb = rgb(t.text)
        chart.has_title = bool(c["title"])
        if c["title"]:
            chart.chart_title.text_frame.text = c["title"]
            r = chart.chart_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.color.rgb = Pt(14), True, rgb(t.text)
        multi = len(c["series"]) > 1 or c["type"] == "pie"
        chart.has_legend = multi
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.color.rgb = rgb(t.text)
        plot = chart.plots[0]
        if c["type"] == "pie":
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(t.chart[i % len(t.chart)])
            plot.has_data_labels = True
            plot.data_labels.show_percentage, plot.data_labels.show_value = True, False
            plot.data_labels.number_format, plot.data_labels.number_format_is_linked = "0%", False
            plot.data_labels.font.color.rgb = rgb("FFFFFF")
            return
        for i, ser in enumerate(plot.series):
            fmt = ser.format
            if c["type"] == "line":
                fmt.line.color.rgb = rgb(t.chart[i % len(t.chart)])
                fmt.line.width = Pt(3)
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = rgb(t.chart[i % len(t.chart)])
        if c["type"] in ("column", "bar"):
            plot.gap_width = 60
            if len(c["series"]) == 1:
                plot.has_data_labels = True
                plot.data_labels.font.color.rgb = rgb(t.text)
                plot.data_labels.font.size = Pt(11)
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        for ax in (chart.category_axis, chart.value_axis):
            ax.tick_labels.font.color.rgb = rgb(t.muted)
            ax.format.line.color.rgb = rgb(t.muted)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = rgb(t.surface if t.dark else "E5E7EB")

    def table_slide(self, sl: dict):
        t, s = self.theme, self._new("table")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        tb = sl["table"]
        top, bottom = (1.95 if sl["subtitle"] else 1.6), (6.45 if cited else 6.8)
        rows, cols = len(tb["rows"]) + 1, len(tb["columns"])
        width = W - 2 * MX
        avail = bottom - top - (0.4 if tb["caption"] else 0)
        size = 16 if rows <= 6 else 14 if rows <= 9 else 12
        col_w = width / cols
        # shrink the type until the tallest row of each column wraps into at most 3 lines and the table fits the space
        while size > 10:
            row_h = [max(_lines(cell, t.body_font, size, (col_w - 0.2) * 72) for cell in r) for r in [tb["columns"]] + tb["rows"]]
            need = sum(n * size * 1.25 / 72 + 0.16 for n in row_h)
            if need <= avail and max(row_h) <= 3:
                break
            size -= 1
        self.fit_report.append({"slide": self.slide_no, "box": "table", "size": size,
                                "fits": sum(n * size * 1.25 / 72 + 0.16 for n in row_h) <= avail})
        gf = s.shapes.add_table(rows, cols, Inches(MX), Inches(top), Inches(width), Inches(min(avail, rows * 0.5)))
        table = gf.table
        for j in range(cols):
            table.columns[j].width = Emu(int(Inches(width) / cols))
        for i, row in enumerate([tb["columns"]] + tb["rows"]):
            for j, val in enumerate(row):
                cell = table.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(t.accent if i == 0 else (t.surface if i % 2 else t.bg))
                cell.margin_left = cell.margin_right = Inches(0.1)
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = ""
                self._runs(tf.paragraphs[0], val, t.body_font, size, "FFFFFF" if i == 0 else t.text, i == 0, False)
        # no entrance animation on tables: OfficeCLI attaches animations to shapes and charts only (measured)
        if tb["caption"]:
            self._text(s, MX, bottom - 0.4, width, 0.4, [(tb["caption"], -1)], size=12, floor=9, color=t.muted, italic=True, name="caption")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def quote(self, sl: dict):
        t, s = self.theme, self._new("quote")
        self._text(s, MX, 0.9, 1.5, 1.6, [("“", -1)], family="Georgia", size=120, floor=80, color=t.accent, name="mark")
        q = self._text(s, MX + 1.2, 1.9, W - 2 * MX - 1.6, 3.4, [(sl["quote"] or sl["title"], -1)], family=t.heading_font,
                       size=34, floor=20, italic=True, anchor=MSO_ANCHOR.MIDDLE, name="quote", spacing=1.15)
        self._animate(q, "fade")
        if sl["attribution"]:
            self._text(s, MX + 1.2, 5.5, W - 2 * MX - 1.6, 0.6, [("— " + sl["attribution"], -1)], size=18, floor=12,
                       color=t.muted, name="attribution")
        cited = self._cited(sl)
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def timeline(self, sl: dict):
        t, s = self.theme, self._new("timeline")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        items = sl["items"]
        n, y = len(items), 4.0
        self._rect(s, MX, y - 0.03, W - 2 * MX, 0.06, t.muted)
        step = (W - 2 * MX) / n
        for i, it in enumerate(items):
            cx = MX + step * (i + 0.5)
            dot = self._rect(s, cx - 0.17, y - 0.17, 0.34, 0.34, t.accent if i % 2 == 0 else t.accent2, MSO_SHAPE.OVAL)
            up = i % 2 == 0
            lab = self._text(s, cx - step / 2 + 0.05, (y - 1.05) if up else (y + 0.3), step - 0.1, 0.7, [(it["label"], -1)],
                             family=t.heading_font, size=20, floor=12, bold=True, color=t.accent if up else t.accent2,
                             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP, name=f"label{i}")
            txt = self._text(s, cx - step / 2 + 0.05, (y - 2.55) if up else (y + 0.95), step - 0.1, 1.45, [(it["text"], -1)],
                             size=14, floor=10, color=t.text, align=PP_ALIGN.CENTER,
                             anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP, name=f"text{i}")
            for shp in (dot, lab, txt) if n <= 6 else (lab,):
                self._animate(shp, "fade")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def stats(self, sl: dict):
        t, s = self.theme, self._new("stats")
        self._title(s, sl["title"], sl["subtitle"])
        cited = self._cited(sl)
        items = sl["items"][:4]
        n = len(items)
        gap = 0.35
        cw = (W - 2 * MX - gap * (n - 1)) / n
        for i, it in enumerate(items):
            x = MX + i * (cw + gap)
            card = self._rect(s, x, 2.3, cw, 3.6, t.surface, MSO_SHAPE.ROUNDED_RECTANGLE)
            card.adjustments[0] = 0.06
            self._rect(s, x, 2.3, cw, 0.09, t.accent if i % 2 == 0 else t.accent2)
            val = self._text(s, x + 0.2, 2.6, cw - 0.4, 1.4, [(it["label"], -1)], family=t.heading_font, size=54, floor=24,
                             bold=True, color=t.accent if i % 2 == 0 else t.accent2, align=PP_ALIGN.CENTER,
                             anchor=MSO_ANCHOR.MIDDLE, name=f"value{i}")
            self._text(s, x + 0.2, 4.1, cw - 0.4, 1.6, [(it["text"], -1)], size=17, floor=11, align=PP_ALIGN.CENTER, name=f"stat{i}")
            self._animate(card, "zoom")
            self._animate(val, "fade")
        self._footer(s, cited)
        self._notes(s, sl, cited)

    def closing(self, sl: dict):
        t, s = self.theme, self._new("closing")
        self._rect(s, -1.5, 3.8, 6, 6, t.accent, MSO_SHAPE.OVAL, alpha=15)
        main = self._text(s, MX, 2.3, W - 2 * MX, 1.6, [(sl["title"] or "Thank you", -1)], family=t.heading_font, size=60,
                          floor=32, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="title")
        self._animate(main, "zoom")
        if sl["subtitle"]:
            self._text(s, MX, 4.0, W - 2 * MX, 0.8, [(sl["subtitle"], -1)], size=24, floor=14, color=t.muted,
                       align=PP_ALIGN.CENTER, name="subtitle")
        self._notes(s, sl, [])

    def sources_slides(self):
        srcs = self.spec["sources"]
        for start in range(0, len(srcs), 8):
            chunk = srcs[start:start + 8]
            s = self._new("sources")
            self._title(s, "Sources" if start == 0 else "Sources (continued)")
            paras = []
            for x in chunk:
                bits = [x["title"], x["publisher"], x["date"]]
                paras.append((f"[{self.src_num[x['id']]}] " + ", ".join(b for b in bits if b) + (f" — {x['url']}" if x["url"] else ""), -1))
            box = self._text(s, MX + 0.25, 1.6, W - 2 * MX - 0.25, 5.1, paras, size=15, floor=9, color=self.theme.text,
                             name="sources", gap=0.6)
            for i, x in enumerate(chunk):   # make each URL a real link
                if x["url"] and i < len(box.text_frame.paragraphs):
                    for r in box.text_frame.paragraphs[i].runs:
                        r.hyperlink.address = x["url"]
            self._footer(s, [])

    # ---------- whole deck ----------

    def build(self) -> Presentation:
        section_no = 0
        slides = self.spec["slides"]
        if slides[0]["layout"] != "title":
            self.title_slide({"title": self.spec["title"], "subtitle": self.spec["subtitle"], "notes": "", "layout": "title"})
        for sl in slides:
            layout = sl["layout"]
            if layout == "title":
                self.title_slide(sl)
            elif layout == "section":
                section_no += 1
                self.section(sl, section_no)
            elif layout == "closing":
                if self.spec["sources"]:
                    self.sources_slides()
                self.closing(sl)
            else:
                getattr(self, {"bullets": "bullets", "two_column": "two_column", "chart": "chart_slide",
                               "table": "table_slide", "quote": "quote", "timeline": "timeline", "stats": "stats"}[layout])(sl)
        if self.spec["sources"] and slides[-1]["layout"] != "closing":
            self.sources_slides()
        return self.prs


def _fix_axis_ids(chart) -> None:
    """python-pptx writes random negative axis ids; the schema wants 0..2147483647 (OfficeCLI validate flags both
    negative values and values above that). The same mask on axId and crossAx keeps the axis pairing intact."""
    for el in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
        el.set("val", str(int(el.get("val")) & 0x7FFFFFFF))


def build(spec: dict, path: str) -> dict:
    b = Builder(spec)
    prs = b.build()
    prs.core_properties.title = spec["title"]
    prs.core_properties.author = spec["author"] or "Neko"
    prs.save(path)
    return {"theme": b.theme.name, "slides": len(prs.slides), "fit": b.fit_report, "animations": b.anim,
            "transition": b.theme.transition}
